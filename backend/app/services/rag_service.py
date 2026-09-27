import io
from typing import List, Dict, Any, Optional
from sqlalchemy.orm import Session
from app.models.research import ResearchDocument, DocumentChunk

class RAGService:
    @staticmethod
    def process_and_store_document(
        db: Session,
        filename: str,
        file_bytes: bytes,
        title: str,
        topic: str = "General",
        region: str = "Global"
    ) -> ResearchDocument:
        full_text = ""
        file_type = "TXT"
        name = filename.lower()

        if name.endswith(".pdf"):
            file_type = "PDF"
            try:
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                for page in reader.pages:
                    full_text += (page.extract_text() or "") + "\n"
            except Exception:
                full_text = file_bytes.decode("utf-8", errors="ignore")
        elif name.endswith(".docx"):
            file_type = "DOCX"
            try:
                import docx
                doc = docx.Document(io.BytesIO(file_bytes))
                for p in doc.paragraphs:
                    full_text += p.text + "\n"
            except Exception:
                full_text = file_bytes.decode("utf-8", errors="ignore")
        else:
            full_text = file_bytes.decode("utf-8", errors="ignore")

        # Chunk text
        chunks = [full_text[i:i+1000] for i in range(0, len(full_text), 800)]
        if not chunks:
            chunks = [full_text]

        doc_record = ResearchDocument(
            filename=filename,
            file_type=file_type,
            title=title,
            topic=topic,
            region=region,
            summary=full_text[:300] + "...",
            full_text=full_text,
            chunks_count=len(chunks)
        )
        db.add(doc_record)
        db.commit()
        db.refresh(doc_record)

        for idx, chunk in enumerate(chunks):
            chunk_rec = DocumentChunk(
                document_id=doc_record.id,
                chunk_index=idx,
                content=chunk,
                keywords=[w for w in chunk.split()[:10] if len(w) > 4]
            )
            db.add(chunk_rec)
        db.commit()

        return doc_record

    @staticmethod
    def search_library(db: Session, query: str, limit: int = 10) -> List[Dict[str, Any]]:
        results = []
        chunks = db.query(DocumentChunk).filter(DocumentChunk.content.ilike(f"%{query}%")).limit(limit).all()
        for c in chunks:
            doc = db.query(ResearchDocument).filter(ResearchDocument.id == c.document_id).first()
            results.append({
                "document_title": doc.title if doc else "Archival Record",
                "filename": doc.filename if doc else "doc.txt",
                "topic": doc.topic if doc else "UNSC",
                "snippet": c.content[:250] + "..."
            })
        return results
